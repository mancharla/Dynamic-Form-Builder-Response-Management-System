from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models.form import Form
from app.models.form_field import FormField
from app.models.form_response import FormResponse


def export_form_responses_to_excel(
    db: Session,
    form_id: int,
) -> tuple[BytesIO, str] | None:

    statement = (
        select(Form)
        .options(
            selectinload(Form.fields),
            selectinload(Form.responses).selectinload(
                FormResponse.details
            ),
        )
        .where(Form.id == form_id)
    )

    form = db.scalar(statement)

    if not form:
        return None

    fields = sorted(
        form.fields,
        key=lambda field: (field.field_order, field.id),
    )

    responses = sorted(
        form.responses,
        key=lambda response: response.submitted_at,
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Responses"

    # ---------------------------------------
    # Header row
    # ---------------------------------------

    headers = [
        "Response ID",
        "User ID",
        "Status",
        "Submitted At",
    ]

    headers.extend(
        field.label
        for field in fields
    )

    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(
            row=1,
            column=column,
            value=header,
        )

        cell.font = Font(bold=True)

    # ---------------------------------------
    # Response rows
    # ---------------------------------------

    for row_number, response in enumerate(
        responses,
        start=2,
    ):
        detail_map = {
            detail.field_id: (
                detail.structured_value
                if detail.structured_value is not None
                else detail.value
            )
            for detail in response.details
        }

        row = [
            response.id,
            response.user_id,
            response.status,
            response.submitted_at.strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
        ]

        for field in fields:
            value = detail_map.get(field.id)

            if isinstance(value, list):
                value = ", ".join(
                    str(item)
                    for item in value
                )

            elif isinstance(value, dict):
                value = str(value)

            row.append(value)

        for column, value in enumerate(
            row,
            start=1,
        ):
            worksheet.cell(
                row=row_number,
                column=column,
                value=value,
            )

    # ---------------------------------------
    # Auto-size columns
    # ---------------------------------------

    for column_cells in worksheet.columns:
        max_length = 0

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)

            max_length = max(
                max_length,
                len(value),
            )

        worksheet.column_dimensions[
            column_cells[0].column_letter
        ].width = min(
            max_length + 2,
            50,
        )

    # ---------------------------------------
    # Create in-memory Excel file
    # ---------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    filename = (
        f"{form.slug}_responses.xlsx"
    )

    return output, filename